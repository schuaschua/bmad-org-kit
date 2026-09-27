#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Build a cost estimate workbook (.xlsx) from an estimate model, and check it for drift.

The model (<name>.estimate.json, in the org config's cost_estimates_folder, default
docs/cost_estimates/) lists every environment's cloud line items, the delivery team, other one-off
costs and the gaps, each line citing where it came from. The workbook has:

  Overall          inputs (days of delivery, working days per month, hours per month, months of run,
                   contingency, exchange rate), the delivery team (employees x days x allocation =
                   mandays, priced from the rate card) and the cost summary gathering every sheet
  <environment>    one sheet per environment: component, Azure service, SKU, unit price, quantity,
                   usage per month, monthly and annual cost, citation and price source
  Rate Card        role -> day rate, from the project's rate card
  Other Costs      one-off costs (licences, pen test, training ...)
  Sources & Gaps   price basis, assumptions, what is not priced yet, and the stamp of every cited file

Every data range is a named Excel table (tblInputs, tblTeam, tblCostSummary, tblEnv<name>, tblRateCard,
tblOtherCosts, and with close_out tblClaudeUsage, tblClaudeUsageDetail, tblClaudePrices, tblTimeline),
so Power BI and Excel Online can load them by name. Formulas get their values when Excel first opens
the file (Power BI reads the values Excel last saved), so <name>.facts.csv is also written beside the
workbook, with every cost line as a plain value for dashboards that read files directly.

Every total is an Excel formula, so changing a blue input cell recalculates the workbook. The
figures are also computed here and printed, so they can be reported without opening Excel.

Model (JSON):
  {"name": "burdeebur", "title": "...", "currency": "AED",           reporting currency
   "cloud_currency": "USD", "fx_rate": 3.6725, "fx_source": "...",    only when prices differ
   "pricing_region": "uaenorth", "price_date": "2026-09-28",
   "assumptions": {"delivery_days": 60, "working_days_per_month": 21, "hours_per_month": 730,
                   "run_months": 12, "contingency_pct": 10},
   "environments": [{"name": "Dev", "delivery_months": null, "run_months": null, "items": [
       {"component": "API", "service": "Azure Kubernetes Service", "sku": "Standard_D4s_v5 node",
        "unit": "1 Hour", "unit_price": 0.235, "quantity": 2, "usage": "hours",
        "cite": "_bmad-output/planning-artifacts/architecture/x/ARCHITECTURE-SPINE.md AD-4",
        "price_source": "Azure Retail Prices API ... meterId ..., fetched 2026-09-28", "note": ""}]}],
   "team": [{"role": "Developer", "employees": 3, "days": null, "allocation": 1, "cite": "..."}],
   "other_costs": [{"item": "Penetration test", "cost": 40000, "cite": "owner: Tech Lead, 2026-09-28"}],
   "gaps": ["Prod DR region: the spine does not say whether there is one"],
   "close_out": {"fx_rate": 3.6725, "fx_source": "...", "in_totals": true}}    optional, at the project's end

  close_out adds the Claude Usage, Claude Usage Detail, Claude Prices and Project Timeline sheets from
  the token ledger's claude_usage.csv and project_timeline.csv (token_usage_folder) and
  _bmad/memory/agent-scrooge/claude-prices.json (paths overridable with "ledger", "prices",
  "timeline"), and a Claude API usage line in the cost summary ("in_totals": false shows it as a memo
  below the grand total instead, for Claude Code on a subscription). fx_rate turns USD into the
  reporting currency.

  cite: "<project-relative file> [<text that appears in it, e.g. AD-4 or a heading>]",
        "owner: <who, date>" for a figure the owner gave, or "assumption: <what and why>".
  usage: units per month per quantity; "hours" means the Hours per month input (always-on).
  null delivery_months / run_months / days mean the Overall inputs.

Rate card: <project>/_bmad/memory/agent-scrooge/rate-card.json (or --rate-card, or "rate_card" in
the model): {"currency": "AED", "updated": "2026-09-28", "source": "...",
             "roles": [{"role": "Developer", "day_rate": 2500, "note": ""}]}

Usage:
    uv run build_estimate.py <project-root> <model.json> [--rate-card FILE] [--out FILE] [--force]
    uv run build_estimate.py <project-root> <model.json> --check     validate only
    uv run build_estimate.py <project-root> <model.json> --drift     cited files changed since the build?

Building writes <name>.xlsx beside the model (or --out) and a "stamp" into the model (build time,
sha256 of every cited file and the rate card, the workbook's save time). It refuses to overwrite a workbook edited after the
last build unless --force. Prints JSON. Exit codes: 0 ok, 1 invalid model or drift found, 2 error.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import tomllib
from datetime import datetime
from pathlib import Path

DATA_DIR = Path("_bmad") / "memory" / "agent-scrooge"
RESERVED = {"overall", "rate card", "other costs", "sources & gaps"}
SPARE_ROWS = 3
DEFAULTS = {"working_days_per_month": 21, "hours_per_month": 730, "run_months": 12, "contingency_pct": 0}

# --------------------------------------------------------------------------- inputs


def org_config(root: Path) -> dict:
    """The [modules.org] section of _bmad/config.toml, with _bmad/custom overrides; {} without one."""
    merged: dict = {}
    for path in (root / "_bmad" / "config.toml", root / "_bmad" / "custom" / "config.toml",
                 root / "_bmad" / "custom" / "config.user.toml"):
        if path.is_file():
            with open(path, "rb") as fh:
                merged.update(tomllib.load(fh).get("modules", {}).get("org", {}))
    return merged


def token_usage_folder(root: Path) -> Path:
    value = str(org_config(root).get("token_usage_folder") or "docs/token_usage").replace("{project-root}", str(root))
    path = Path(value)
    return path if path.is_absolute() else root / path


def estimates_folder(root: Path) -> Path:
    value = str(org_config(root).get("cost_estimates_folder") or "docs/cost_estimates")
    value = value.replace("{project-root}", str(root))
    path = Path(value)
    return path if path.is_absolute() else root / path


def load_rate_card(root: Path, model: dict, override: Path | None) -> tuple[dict, Path | None]:
    if "rate_card" in model:
        return model["rate_card"], None
    path = override or root / DATA_DIR / "rate-card.json"
    if not path.is_file():
        return {}, path
    return json.loads(path.read_text(encoding="utf-8")), path


def is_number(value, minimum: float = 0, strict: bool = False) -> bool:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return False
    return value > minimum if strict else value >= minimum


def split_cite(cite: str) -> tuple[str, str, str]:
    """(kind, file, ref): kind is file, owner or assumption."""
    cite = cite.strip()
    for kind in ("owner", "assumption"):
        if cite.lower().startswith(kind + ":"):
            return kind, "", cite[len(kind) + 1:].strip()
    file, _, ref = cite.partition(" ")
    return "file", file, ref.strip()


def check_cite(root: Path, cite, where: str, errors: list[str]) -> None:
    if not isinstance(cite, str) or not cite.strip():
        errors.append(f"{where}: no cite (a project file, 'owner: <who, date>' or 'assumption: <why>')")
        return
    kind, file, ref = split_cite(cite)
    if kind != "file":
        if not ref:
            errors.append(f"{where}: '{kind}:' cite says nothing")
        return
    path = (root / file).resolve()
    if not path.is_relative_to(root.resolve()):
        errors.append(f"{where}: cited file is outside the project: {file}")
    elif not path.is_file():
        errors.append(f"{where}: cited file does not exist: {file}")
    elif ref and ref.lower() not in path.read_text(encoding="utf-8", errors="replace").lower():
        errors.append(f"{where}: '{ref}' does not appear in {file}")


def validate(root: Path, model: dict, card: dict) -> list[str]:
    errors: list[str] = []
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", str(model.get("name", ""))):
        errors.append("name: letters, digits, '.', '_' and '-' only")
    ccy = model.get("currency")
    if not (isinstance(ccy, str) and re.fullmatch(r"[A-Z]{3}", ccy)):
        errors.append("currency: a three-letter code such as USD or AED")
    cloud_ccy = model.get("cloud_currency", ccy)
    if cloud_ccy != ccy:
        if not is_number(model.get("fx_rate"), strict=True):
            errors.append(f"fx_rate: needed to turn {cloud_ccy} prices into {ccy}")
        if not model.get("fx_source"):
            errors.append("fx_source: say where the exchange rate came from and when")
    a = model.get("assumptions", {})
    if not is_number(a.get("delivery_days"), strict=True):
        errors.append("assumptions.delivery_days: the number of working days of delivery")
    for key in ("working_days_per_month", "hours_per_month"):
        if key in a and not is_number(a[key], strict=True):
            errors.append(f"assumptions.{key}: a positive number")
    for key in ("run_months", "contingency_pct"):
        if key in a and not is_number(a[key]):
            errors.append(f"assumptions.{key}: zero or more")

    envs = model.get("environments") or []
    if not envs:
        errors.append("environments: at least one")
    seen: set[str] = set()
    for e_i, env in enumerate(envs):
        name = str(env.get("name", ""))
        where = f"environments[{e_i}] '{name}'"
        if not name or len(name) > 31 or re.search(r"[\[\]:*?/\\']", name):
            errors.append(f"{where}: sheet name must be 1-31 characters without []:*?/\\'")
        if name.lower() in RESERVED or name.lower() in seen:
            errors.append(f"{where}: name clashes with another sheet")
        seen.add(name.lower())
        for key in ("delivery_months", "run_months"):
            if env.get(key) is not None and not is_number(env[key]):
                errors.append(f"{where}.{key}: zero or more, or null for the Overall input")
        if not env.get("items"):
            errors.append(f"{where}: no items (an environment with nothing priced belongs in gaps)")
        for i_i, item in enumerate(env.get("items") or []):
            iw = f"{name} item {i_i + 1} '{item.get('component', '')}'"
            for key in ("component", "service", "unit", "price_source"):
                if not item.get(key):
                    errors.append(f"{iw}: {key} missing")
            if not is_number(item.get("unit_price")):
                errors.append(f"{iw}: unit_price must be a number (an unpriced item belongs in gaps)")
            if not is_number(item.get("quantity"), strict=True):
                errors.append(f"{iw}: quantity must be above zero")
            usage = item.get("usage", "hours" if "hour" in str(item.get("unit", "")).lower() else None)
            if usage != "hours" and not is_number(usage, strict=True):
                errors.append(f"{iw}: usage must be 'hours' or a number above zero")
            check_cite(root, item.get("cite"), iw, errors)

    if card.get("currency") and ccy and card["currency"] != ccy:
        errors.append(f"rate card is in {card['currency']}, the estimate in {ccy}")
    rates = {}
    for r in card.get("roles", []):
        if not r.get("role") or not is_number(r.get("day_rate"), strict=True):
            errors.append(f"rate card: role '{r.get('role', '')}' needs a day_rate above zero")
        rates[str(r.get("role", "")).lower()] = r.get("day_rate")
    for t_i, member in enumerate(model.get("team") or []):
        tw = f"team[{t_i}] '{member.get('role', '')}'"
        if str(member.get("role", "")).lower() not in rates:
            errors.append(f"{tw}: role is not on the rate card")
        if not is_number(member.get("employees"), strict=True):
            errors.append(f"{tw}: employees must be above zero")
        if member.get("days") is not None and not is_number(member["days"], strict=True):
            errors.append(f"{tw}: days must be above zero, or null for the days of delivery")
        alloc = member.get("allocation", 1)
        if not is_number(alloc, strict=True) or alloc > 1:
            errors.append(f"{tw}: allocation is a fraction above 0 and at most 1")
        check_cite(root, member.get("cite"), tw, errors)
    for o_i, other in enumerate(model.get("other_costs") or []):
        ow = f"other_costs[{o_i}] '{other.get('item', '')}'"
        if not other.get("item") or not is_number(other.get("cost")):
            errors.append(f"{ow}: needs item and a cost of zero or more")
        check_cite(root, other.get("cite"), ow, errors)
    if not isinstance(model.get("gaps", []), list):
        errors.append("gaps: a list of strings")
    return errors


# --------------------------------------------------------------------------- close-out (Claude usage, timeline)


def pricing_id(model: str) -> str:
    """A model id as the logs write it -> the pricing page's id: 'claude-haiku-4-5-20251001' -> 'claude-haiku-4-5'."""
    return re.sub(r"-\d{8}$", "", re.sub(r"\[.*?\]", "", model.strip().lower()))


def load_close_out(root: Path, model: dict) -> tuple[dict | None, list[str]]:
    """The Claude usage, Claude prices and project timeline the close-out sheets are built from."""
    spec = model.get("close_out")
    if spec is None:
        return None, []
    errors: list[str] = []
    ledger_dir = token_usage_folder(root)

    def path(key: str, default: Path) -> Path:
        value = spec.get(key)
        return (root / value) if value else default

    ledger, prices_file = path("ledger", ledger_dir / "claude_usage.csv"), path("prices", root / DATA_DIR / "claude-prices.json")
    timeline_file = path("timeline", ledger_dir / "project_timeline.csv")
    usage, prices, timeline = [], {}, []
    if not ledger.is_file():
        errors.append(f"close_out: no Claude usage ledger at {ledger}; refresh it with token_report.py")
    else:
        with ledger.open(newline="", encoding="utf-8") as fh:
            usage = [r for r in csv.DictReader(fh) if int(r.get("calls") or 0)]
    if not prices_file.is_file():
        errors.append(f"close_out: no Claude prices at {prices_file}; read them with claude_prices.py --save")
    else:
        prices = json.loads(prices_file.read_text(encoding="utf-8"))
    if timeline_file.is_file():
        with timeline_file.open(newline="", encoding="utf-8") as fh:
            timeline = list(csv.DictReader(fh))
    ccy = model.get("currency")
    if ccy != "USD" and not is_number(spec.get("fx_rate"), strict=True):
        errors.append(f"close_out.fx_rate: Claude prices are in USD; give the rate to {ccy} (and fx_source)")
    known = prices.get("models", {})
    tokens = ("input_tokens", "cache_write_5m_tokens", "cache_write_1h_tokens", "cache_read_tokens", "output_tokens")
    unpriced = sorted({r["model"] for r in usage if pricing_id(r["model"]) not in known
                       and any(int(r.get(t) or 0) for t in tokens)})
    return {"usage": usage, "prices": prices, "timeline": timeline, "unpriced": unpriced,
            "fx_rate": spec.get("fx_rate", 1) if ccy != "USD" else 1, "fx_source": spec.get("fx_source", ""),
            "in_totals": spec.get("in_totals", True),
            "files": [p for p in (ledger, prices_file, timeline_file) if p.is_file()]}, errors


def claude_cost(close: dict) -> tuple[float, int]:
    """(USD at API list prices, calls)."""
    known = close["prices"].get("models", {})
    usd, calls = 0.0, 0
    for r in close["usage"]:
        calls += int(r["calls"])
        price = known.get(pricing_id(r["model"]))
        if price:
            usd += sum(int(r[f"{k}_tokens"] or 0) * price[k] for k in
                       ("input", "cache_write_5m", "cache_write_1h", "cache_read", "output")) / 1_000_000
    return usd, calls


# --------------------------------------------------------------------------- figures


def settings(model: dict) -> dict:
    a = {**DEFAULTS, **model.get("assumptions", {})}
    a["contingency"] = a["contingency_pct"] / 100
    a["delivery_months"] = a["delivery_days"] / a["working_days_per_month"]
    a["fx_rate"] = model.get("fx_rate", 1) if model.get("cloud_currency", model["currency"]) != model["currency"] else 1
    return a


def item_usage(item: dict, a: dict) -> float:
    usage = item.get("usage", "hours" if "hour" in str(item.get("unit", "")).lower() else None)
    return a["hours_per_month"] if usage == "hours" else usage


def compute(model: dict, card: dict, close: dict | None = None) -> dict:
    a = settings(model)
    rates = {str(r["role"]).lower(): r["day_rate"] for r in card.get("roles", [])}
    envs = []
    for env in model["environments"]:
        monthly = sum(i["unit_price"] * i["quantity"] * item_usage(i, a) * a["fx_rate"] for i in env["items"])
        d_months = a["delivery_months"] if env.get("delivery_months") is None else env["delivery_months"]
        r_months = a["run_months"] if env.get("run_months") is None else env["run_months"]
        envs.append({"name": env["name"], "monthly": monthly, "delivery": monthly * d_months,
                     "run": monthly * r_months})
    mandays = people = 0.0
    for m in model.get("team") or []:
        days = a["delivery_days"] if m.get("days") is None else m["days"]
        md = m["employees"] * days * m.get("allocation", 1)
        mandays += md
        people += md * rates[m["role"].lower()]
    other = sum(o["cost"] for o in model.get("other_costs") or [])
    claude = None
    if close:
        usd, calls = claude_cost(close)
        claude = {"calls": calls, "usd": round(usd, 2), "cost": round(usd * close["fx_rate"], 2),
                  "in_totals": bool(close["in_totals"]), "unpriced_models": close["unpriced"]}
    delivery = sum(e["delivery"] for e in envs) + people + other + (claude["cost"] if claude and claude["in_totals"] else 0)
    run = sum(e["run"] for e in envs)
    r2 = lambda x: round(x, 2)  # noqa: E731
    return {
        "currency": model["currency"],
        "delivery_months": r2(a["delivery_months"]),
        "environments": [{k: (r2(v) if isinstance(v, float | int) else v) for k, v in e.items()} for e in envs],
        "team": {"employees": sum(m["employees"] for m in model.get("team") or []), "mandays": r2(mandays),
                 "cost": r2(people)},
        "other_costs": r2(other),
        **({"claude": claude} if claude else {}),
        "subtotal": {"delivery": r2(delivery), "run": r2(run), "total": r2(delivery + run)},
        "contingency": r2((delivery + run) * a["contingency"]),
        "grand_total": {"delivery": r2(delivery * (1 + a["contingency"])),
                        "run": r2(run * (1 + a["contingency"])),
                        "total": r2((delivery + run) * (1 + a["contingency"]))},
        "gaps": len(model.get("gaps") or []),
    }


# --------------------------------------------------------------------------- stamp and drift


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def cited_files(model: dict) -> list[str]:
    cites = [i.get("cite", "") for e in model.get("environments", []) for i in e.get("items", [])]
    cites += [m.get("cite", "") for m in model.get("team") or []]
    cites += [o.get("cite", "") for o in model.get("other_costs") or []]
    files = {split_cite(c)[1] for c in cites if isinstance(c, str)}
    return sorted(f for f in files if f)


def make_stamp(root: Path, model: dict, card_path: Path | None, extra: list[Path] = ()) -> dict:
    files = {f: sha(root / f) for f in cited_files(model)}
    for path in extra:
        files[path.resolve().relative_to(root.resolve()).as_posix()] = sha(path)
    if card_path and card_path.is_file():
        rel = card_path.resolve().relative_to(root.resolve()).as_posix() if card_path.resolve().is_relative_to(
            root.resolve()) else str(card_path)
        files[rel] = sha(card_path)
    return {"built": datetime.now().isoformat(timespec="seconds"), "files": files}


def drift(root: Path, model: dict) -> dict:
    stamp = model.get("stamp")
    if not stamp:
        return {"status": "never-built", "changed": [], "missing": []}
    changed, missing = [], []
    for rel, digest in stamp.get("files", {}).items():
        path = Path(rel) if Path(rel).is_absolute() else root / rel
        if not path.is_file():
            missing.append(rel)
        elif sha(path) != digest:
            changed.append(rel)
    return {"status": "drift" if changed or missing else "current", "built": stamp.get("built"),
            "changed": changed, "missing": missing}


# --------------------------------------------------------------------------- workbook


def write_workbook(model: dict, card: dict, stamp: dict, out: Path, close: dict | None = None) -> list[str]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.table import Table, TableStyleInfo

    a = settings(model)
    ccy = model["currency"]
    cloud_ccy = model.get("cloud_currency", ccy)
    region = model.get("pricing_region", "")
    price_date = model.get("price_date", "")

    MONEY, QTY, PCT, MONTHS = "#,##0.00", "#,##0.##", "0%", "0.0"
    BLUE, GREEN = Font(color="0000FF"), Font(color="008000")
    BOLD = Font(bold=True)
    HEAD = Font(bold=True, color="FFFFFF")
    HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
    SECTION = Font(bold=True, size=12, color="1F4E78")
    TITLE = Font(bold=True, size=14)
    TOP = Border(top=Side(style="thin"))
    WRAP = Alignment(wrap_text=True, vertical="top")

    def q(sheet: str) -> str:
        return "'" + sheet.replace("'", "''") + "'"

    def header(ws, row: int, labels: list[str]) -> None:
        for c, label in enumerate(labels, 1):
            cell = ws.cell(row=row, column=c, value=label)
            cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, Alignment(wrap_text=True, vertical="center")

    def put(ws, ref: str, value, fmt: str | None = None, font=None):
        cell = ws[ref]
        cell.value = value
        if fmt:
            cell.number_format = fmt
        if font is not None:
            cell.font = font
        return cell

    def widths(ws, values: list[int]) -> None:
        for i, w in enumerate(values, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    tables: set[str] = set()

    def add_table(ws, name: str, first_col: str, header_row: int, last_col: str, last_row: int) -> None:
        """Name a data range as an Excel table so Power BI and Excel Online pick it up by name."""
        name = re.sub(r"[^A-Za-z0-9_]", "", name)
        base, n = name, 2
        while name.lower() in tables:
            name, n = f"{base}{n}", n + 1
        tables.add(name.lower())
        t = Table(displayName=name, ref=f"{first_col}{header_row}:{last_col}{max(last_row, header_row + 1)}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)

    def total_row(ws, row: int, last_col: int) -> None:
        for c in range(1, last_col + 1):
            ws.cell(row=row, column=c).border = TOP
            ws.cell(row=row, column=c).font = BOLD

    wb = Workbook()
    ov = wb.active
    ov.title = "Overall"

    # --- Environments (built first so Overall can point at their totals)
    env_totals: list[tuple[str, int]] = []
    for env in model["environments"]:
        ws = wb.create_sheet(env["name"])
        put(ws, "A1", f"{env['name']} environment", font=TITLE)
        fx_note = f", converted to {ccy} at the Overall exchange rate" if cloud_ccy != ccy else ""
        put(ws, "A2", f"Azure retail list prices (pay-as-you-go) in {region}, {cloud_ccy}, fetched {price_date}"
                      f"{fx_note}. Blue cells are inputs.")
        header(ws, 4, ["#", "Component", "Azure service", "SKU / tier", "Unit", f"Unit price ({cloud_ccy})",
                       "Quantity", "Usage per month", f"Monthly cost ({ccy})", f"Annual cost ({ccy})",
                       "Cited from", "Price source", "Note", "Environment"])
        row = 5
        for n, item in enumerate(env["items"], 1):
            usage = item.get("usage", "hours" if "hour" in str(item.get("unit", "")).lower() else None)
            ws.cell(row=row, column=1, value=n)
            ws.cell(row=row, column=2, value=item["component"])
            ws.cell(row=row, column=3, value=item["service"])
            ws.cell(row=row, column=4, value=item.get("sku", ""))
            ws.cell(row=row, column=5, value=item["unit"])
            put(ws, f"F{row}", item["unit_price"], "#,##0.0000", BLUE)
            put(ws, f"G{row}", item["quantity"], QTY, BLUE)
            if usage == "hours":
                put(ws, f"H{row}", "=HoursPerMonth", QTY)
            else:
                put(ws, f"H{row}", usage, QTY, BLUE)
            put(ws, f"I{row}", f"=F{row}*G{row}*H{row}*FxRate", MONEY)
            put(ws, f"J{row}", f"=I{row}*12", MONEY)
            ws.cell(row=row, column=11, value=item["cite"])
            ws.cell(row=row, column=12, value=item["price_source"])
            ws.cell(row=row, column=13, value=item.get("note", ""))
            ws.cell(row=row, column=14, value=env["name"])
            for c in (11, 12, 13):
                ws.cell(row=row, column=c).alignment = WRAP
            row += 1
        for _ in range(SPARE_ROWS):
            for col in "FGH":
                ws[f"{col}{row}"].font = BLUE
            put(ws, f"I{row}", f'=IF(F{row}="","",F{row}*G{row}*H{row}*FxRate)', MONEY)
            put(ws, f"J{row}", f'=IF(I{row}="","",I{row}*12)', MONEY)
            ws.cell(row=row, column=14, value=env["name"])
            row += 1
        add_table(ws, f"tblEnv{env['name']}", "A", 4, "N", row - 1)
        put(ws, f"A{row}", "Total")
        put(ws, f"I{row}", f"=SUM(I5:I{row - 1})", MONEY)
        put(ws, f"J{row}", f"=SUM(J5:J{row - 1})", MONEY)
        total_row(ws, row, 13)
        env_totals.append((env["name"], row))
        widths(ws, [5, 26, 26, 24, 12, 14, 10, 12, 16, 16, 40, 48, 30, 14])
        ws.freeze_panes = "A5"

    # --- Rate Card
    rc = wb.create_sheet("Rate Card", 1)
    put(rc, "A1", "Employee rate card", font=TITLE)
    put(rc, "A2", f"Day rates in {ccy}. Source: {card.get('source', 'not stated')}; updated "
                  f"{card.get('updated', 'not stated')}. Blue cells are inputs; the Overall team table looks roles up here.")
    header(rc, 4, ["Role", f"Day rate ({ccy})", "Note"])
    r = 5
    for role in card.get("roles", []):
        rc.cell(row=r, column=1, value=role["role"])
        put(rc, f"B{r}", role["day_rate"], MONEY, BLUE)
        rc.cell(row=r, column=3, value=role.get("note", ""))
        r += 1
    add_table(rc, "tblRateCard", "A", 4, "C", r - 1)
    rate_last = max(r + 20, 50)  # room for roles added by hand
    widths(rc, [30, 16, 50])
    rc.freeze_panes = "A5"

    # --- Other Costs
    oc = wb.create_sheet("Other Costs", 2)
    put(oc, "A1", "Other one-off costs", font=TITLE)
    put(oc, "A2", f"Licences, tests, training and other costs outside the environments and the team, in {ccy}.")
    header(oc, 4, ["#", "Item", f"Amount ({ccy})", "Cited from", "Note"])
    r = 5
    for n, other in enumerate(model.get("other_costs") or [], 1):
        oc.cell(row=r, column=1, value=n)
        oc.cell(row=r, column=2, value=other["item"])
        put(oc, f"C{r}", other["cost"], MONEY, BLUE)
        oc.cell(row=r, column=4, value=other["cite"])
        oc.cell(row=r, column=5, value=other.get("note", ""))
        r += 1
    for _ in range(SPARE_ROWS):
        oc[f"C{r}"].font = BLUE
        r += 1
    add_table(oc, "tblOtherCosts", "A", 4, "E", r - 1)
    put(oc, f"A{r}", "Total")
    put(oc, f"C{r}", f"=SUM(C5:C{r - 1})", MONEY)
    total_row(oc, r, 5)
    other_total = r
    widths(oc, [5, 36, 16, 44, 30])

    # --- Overall: inputs
    put(ov, "A1", model.get("title") or f"{model['name']} cost estimate", font=TITLE)
    put(ov, "A2", f"All costs in {ccy}. Cloud: Azure retail list prices in {region} fetched {price_date}, "
                  "no discounts or reservations. Blue cells are inputs; black cells are formulas; green cells link to other sheets.")
    put(ov, "A3", "Inputs", font=SECTION)
    header(ov, 4, ["Input", "Value"])
    inputs = [
        ("DeliveryDays", "Days of delivery (working days)", a["delivery_days"], QTY, True),
        ("WorkingDaysPerMonth", "Working days per month", a["working_days_per_month"], QTY, True),
        ("DeliveryMonths", "Months of delivery", "=DeliveryDays/WorkingDaysPerMonth", MONTHS, False),
        ("HoursPerMonth", "Hours per month (always-on services)", a["hours_per_month"], QTY, True),
        ("RunMonths", "Months of run after go-live", a["run_months"], QTY, True),
        ("Contingency", "Contingency", a["contingency"], PCT, True),
        ("FxRate", f"Exchange rate (1 {cloud_ccy} = x {ccy})", a["fx_rate"], "0.0000", True),
    ]
    if close:
        inputs.append(("ClaudeFxRate", f"Exchange rate for Claude prices (1 USD = x {ccy})", close["fx_rate"], "0.0000", True))
    for offset, (name, label, value, fmt, is_input) in enumerate(inputs):
        r = 5 + offset
        ov.cell(row=r, column=1, value=label)
        put(ov, f"B{r}", value, fmt, BLUE if is_input else None)
        wb.defined_names[name] = DefinedName(name, attr_text=f"Overall!$B${r}")
    add_table(ov, "tblInputs", "A", 4, "B", 4 + len(inputs))
    if cloud_ccy != ccy:
        ov.cell(row=11, column=3, value=model.get("fx_source", ""))
    if close and ccy != "USD":
        ov.cell(row=12, column=3, value=close["fx_source"])

    # --- Overall: delivery team
    r = 5 + len(inputs) + 1
    put(ov, f"A{r}", "Delivery team", font=SECTION)
    r += 1
    header(ov, r, ["Role", f"Day rate ({ccy})", "Employees", "Days on delivery", "Allocation", "Mandays",
                   f"Cost ({ccy})", "Cited from"])
    first = r + 1
    rate_range = f"{q('Rate Card')}!$A$5:$B${rate_last}"
    for m in model.get("team") or []:
        r += 1
        ov.cell(row=r, column=1, value=m["role"])
        put(ov, f"B{r}", f"=VLOOKUP(A{r},{rate_range},2,FALSE)", MONEY, GREEN)
        put(ov, f"C{r}", m["employees"], QTY, BLUE)
        if m.get("days") is None:
            put(ov, f"D{r}", "=DeliveryDays", QTY)
        else:
            put(ov, f"D{r}", m["days"], QTY, BLUE)
        put(ov, f"E{r}", m.get("allocation", 1), PCT, BLUE)
        put(ov, f"F{r}", f"=C{r}*D{r}*E{r}", QTY)
        put(ov, f"G{r}", f"=F{r}*B{r}", MONEY)
        ov.cell(row=r, column=8, value=m["cite"])
    for _ in range(SPARE_ROWS):
        r += 1
        put(ov, f"B{r}", f'=IF(A{r}="","",VLOOKUP(A{r},{rate_range},2,FALSE))', MONEY, GREEN)
        for col in "CDE":
            ov[f"{col}{r}"].font = BLUE
        ov[f"E{r}"].number_format = PCT
        put(ov, f"F{r}", f'=IF(A{r}="","",C{r}*D{r}*E{r})', QTY)
        put(ov, f"G{r}", f'=IF(A{r}="","",F{r}*B{r})', MONEY)
    add_table(ov, "tblTeam", "A", first - 1, "H", r)
    r += 1
    team_total = r
    put(ov, f"A{r}", "Total")
    put(ov, f"C{r}", f"=SUM(C{first}:C{r - 1})", QTY)
    put(ov, f"F{r}", f"=SUM(F{first}:F{r - 1})", QTY)
    put(ov, f"G{r}", f"=SUM(G{first}:G{r - 1})", MONEY)
    total_row(ov, r, 8)

    # --- Overall: cost summary
    r += 2
    put(ov, f"A{r}", "Cost summary", font=SECTION)
    r += 1
    header(ov, r, ["Item", f"Monthly ({ccy})", "Months in delivery", f"Delivery cost ({ccy})",
                   "Months in run", f"Run cost ({ccy})", f"Total ({ccy})", "From"])
    first = r + 1
    for env, (name, env_total) in zip(model["environments"], env_totals):
        r += 1
        ov.cell(row=r, column=1, value=f"{name} environment")
        put(ov, f"B{r}", f"={q(name)}!$I${env_total}", MONEY, GREEN)
        if env.get("delivery_months") is None:
            put(ov, f"C{r}", "=DeliveryMonths", MONTHS)
        else:
            put(ov, f"C{r}", env["delivery_months"], MONTHS, BLUE)
        put(ov, f"D{r}", f"=B{r}*C{r}", MONEY)
        if env.get("run_months") is None:
            put(ov, f"E{r}", "=RunMonths", MONTHS)
        else:
            put(ov, f"E{r}", env["run_months"], MONTHS, BLUE)
        put(ov, f"F{r}", f"=B{r}*E{r}", MONEY)
        put(ov, f"G{r}", f"=D{r}+F{r}", MONEY)
        ov.cell(row=r, column=8, value=f"{name} sheet")
    r += 1
    ov.cell(row=r, column=1, value="Delivery team")
    put(ov, f"D{r}", f"=G{team_total}", MONEY, GREEN)
    put(ov, f"G{r}", f"=D{r}+F{r}", MONEY)
    ov.cell(row=r, column=8, value="Delivery team table above")
    r += 1
    ov.cell(row=r, column=1, value="Other one-off costs")
    put(ov, f"D{r}", f"={q('Other Costs')}!$C${other_total}", MONEY, GREEN)
    put(ov, f"G{r}", f"=D{r}+F{r}", MONEY)
    ov.cell(row=r, column=8, value="Other Costs sheet")
    claude_ref = None
    if close:
        claude_ref = "=" + q("Claude Usage") + "!$I$" + str(5 + len(claude_models(close)))
        if close["in_totals"]:
            r += 1
            ov.cell(row=r, column=1, value="Claude API usage (actual, API list prices)")
            put(ov, f"D{r}", claude_ref, MONEY, GREEN)
            put(ov, f"G{r}", f"=D{r}+F{r}", MONEY)
            ov.cell(row=r, column=8, value="Claude Usage sheet")
    add_table(ov, "tblCostSummary", "A", first - 1, "H", r)
    r += 1
    sub = r
    put(ov, f"A{r}", "Subtotal")
    for col in "DFG":
        put(ov, f"{col}{r}", f"=SUM({col}{first}:{col}{r - 1})", MONEY)
    total_row(ov, r, 8)
    r += 1
    ov.cell(row=r, column=1, value="Contingency")
    for col in "DF":
        put(ov, f"{col}{r}", f"={col}{sub}*Contingency", MONEY)
    put(ov, f"G{r}", f"=D{r}+F{r}", MONEY)
    r += 1
    put(ov, f"A{r}", "Grand total")
    for col in "DFG":
        put(ov, f"{col}{r}", f"={col}{sub}+{col}{r - 1}", MONEY)
    total_row(ov, r, 8)
    for col in "ADFG":
        ov[f"{col}{r}"].font = Font(bold=True, size=12)
    if close and not close["in_totals"]:
        r += 2
        ov.cell(row=r, column=1, value="Memo: Claude API usage at API list prices (not in the totals: Claude Code was on a subscription)")
        put(ov, f"D{r}", claude_ref, MONEY, GREEN)
    widths(ov, [36, 16, 12, 18, 12, 18, 18, 44])

    # --- Sources & Gaps
    sg = wb.create_sheet("Sources & Gaps")
    put(sg, "A1", "Sources, assumptions and gaps", font=TITLE)
    lines: list[tuple[str, str]] = [("Price basis", ""),
                                    ("", f"Azure Retail Prices API list prices (pay-as-you-go, no discounts), region {region}, "
                                         f"{cloud_ccy}, fetched {price_date}.")]
    if cloud_ccy != ccy:
        lines.append(("", f"Exchange rate 1 {cloud_ccy} = {a['fx_rate']} {ccy}: {model.get('fx_source', '')}"))
    lines.append(("", f"Rate card: {card.get('source', 'not stated')}, updated {card.get('updated', 'not stated')}."))
    if close:
        pr = close["prices"]
        lines.append(("", f"Claude API list prices from {pr.get('source', '')}, read {pr.get('fetched', '')}, USD per million "
                          "tokens. Claude Code on a subscription is billed per seat: this is the API-equivalent cost."))
        if close["unpriced"]:
            lines.append(("", "Claude models with no price on the page (their tokens are not costed): " + ", ".join(close["unpriced"])))
    lines += [("", ""), ("Gaps: not priced and not in any total", "")]
    lines += [("", g) for g in model.get("gaps") or []] or [("", "None recorded.")]
    assumed = []
    for env in model["environments"]:
        for item in env["items"]:
            kind, _, ref = split_cite(item["cite"])
            if kind != "file":
                assumed.append(f"{env['name']} / {item['component']}: {kind}: {ref}")
    for m in model.get("team") or []:
        kind, _, ref = split_cite(m["cite"])
        if kind != "file":
            assumed.append(f"Team / {m['role']}: {kind}: {ref}")
    for o in model.get("other_costs") or []:
        kind, _, ref = split_cite(o["cite"])
        if kind != "file":
            assumed.append(f"Other / {o['item']}: {kind}: {ref}")
    lines += [("", ""), ("Assumptions and owner statements (not from the BMad artifacts)", "")]
    lines += [("", t) for t in assumed] or [("", "None.")]
    lines += [("", ""), (f"Stamp: built {stamp['built']}; sha256 of each cited file", "")]
    lines += [("", f"{path}  {digest[:12]}") for path, digest in stamp["files"].items()]
    for i, (head, text) in enumerate(lines, 3):
        if head:
            put(sg, f"A{i}", head, font=SECTION)
        if text:
            sg.cell(row=i, column=2, value=text).alignment = Alignment(wrap_text=False)
    widths(sg, [4, 120])

    if close:
        write_close_out(wb, close, ccy, dict(put=put, header=header, widths=widths, total_row=total_row,
                                             add_table=add_table, q=q, TITLE=TITLE, SECTION=SECTION, MONEY=MONEY,
                                             QTY=QTY, BLUE=BLUE, GREEN=GREEN, BOLD=BOLD))
    wb.move_sheet("Sources & Gaps", offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Sources & Gaps"))
    wb.calculation.fullCalcOnLoad = True
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return wb.sheetnames


def claude_models(close: dict) -> list[str]:
    return sorted({r["model"] for r in close["usage"]})


def iso_datetime(value: str):
    """'2026-09-27T11:02:59.394Z' or '2026-09-27' -> a naive datetime Excel and Power BI read as a date."""
    if not value:
        return None
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        return value


def write_close_out(wb, close: dict, ccy: str, k: dict) -> None:
    put, header, widths, add_table, q = k["put"], k["header"], k["widths"], k["add_table"], k["q"]
    MONEY, QTY, BLUE = k["MONEY"], k["QTY"], k["BLUE"]
    DATE, TOKENS = "yyyy-mm-dd hh:mm", "#,##0"
    kinds = ("input", "cache_write_5m", "cache_write_1h", "cache_read", "output")

    # Claude Prices
    pr = close["prices"]
    ps = wb.create_sheet("Claude Prices")
    put(ps, "A1", "Claude API list prices", font=k["TITLE"])
    put(ps, "A2", f"USD per million tokens, from {pr.get('source', '')}, read {pr.get('fetched', '')}. Blue cells are inputs.")
    header(ps, 4, ["Pricing model", "Name", "Input", "Cache write 5m", "Cache write 1h", "Cache read", "Output"])
    r = 4
    for mid, price in sorted(pr.get("models", {}).items()):
        r += 1
        ps.cell(row=r, column=1, value=mid)
        ps.cell(row=r, column=2, value=price.get("name", mid))
        for c, key in enumerate(kinds, 3):
            put(ps, f"{chr(64 + c)}{r}", price[key], "#,##0.00", BLUE)
    add_table(ps, "tblClaudePrices", "A", 4, "G", r)
    price_range = f"{q('Claude Prices')}!$A$5:$G${max(r, 5)}"
    widths(ps, [24, 22, 12, 14, 14, 12, 12])

    # Claude Usage Detail: one row per work item x model
    ds = wb.create_sheet("Claude Usage Detail")
    put(ds, "A1", "Claude API calls by work item and model", font=k["TITLE"])
    put(ds, "A2", "From the token ledger (claude_usage.csv). Cost = tokens x the Claude Prices sheet; models with no price cost 0.")
    header(ds, 4, ["Kind", "Work item", "Story no", "Title", "Model", "Pricing model", "First call", "Last call", "Calls",
                   "Input tokens", "Cache write 5m tokens", "Cache write 1h tokens", "Cache read tokens", "Output tokens",
                   "Cost (USD)", f"Cost ({ccy})"])
    r = 4
    for u in close["usage"]:
        r += 1
        values = [u["kind"], u["work_item"], u.get("story_no", ""), u.get("title", ""), u["model"], pricing_id(u["model"]),
                  iso_datetime(u.get("first_call", "")), iso_datetime(u.get("last_call", "")), int(u["calls"])]
        values += [int(u.get(f"{key}_tokens") or 0) for key in kinds]
        for c, v in enumerate(values, 1):
            cell = ds.cell(row=r, column=c, value=v)
            cell.number_format = DATE if c in (7, 8) else TOKENS if c >= 9 else "General"
        vl = [f"J{r}*VLOOKUP($F{r},{price_range},3,FALSE)", f"K{r}*VLOOKUP($F{r},{price_range},4,FALSE)",
              f"L{r}*VLOOKUP($F{r},{price_range},5,FALSE)", f"M{r}*VLOOKUP($F{r},{price_range},6,FALSE)",
              f"N{r}*VLOOKUP($F{r},{price_range},7,FALSE)"]
        put(ds, f"O{r}", f"=IFERROR(({'+'.join(vl)})/1000000,0)", MONEY)
        put(ds, f"P{r}", f"=O{r}*ClaudeFxRate", MONEY)
    add_table(ds, "tblClaudeUsageDetail", "A", 4, "P", r)
    last_detail = max(r, 5)
    widths(ds, [10, 30, 9, 30, 26, 22, 17, 17, 9, 14, 14, 14, 16, 14, 14, 14])
    ds.freeze_panes = "A5"

    # Claude Usage: by model, totals
    us = wb.create_sheet("Claude Usage")
    put(us, "A1", "Claude API usage", font=k["TITLE"])
    put(us, "A2", "API calls and tokens counted from this project's Claude Code logs, costed at Claude API list prices. "
                  "On a Pro, Max, Team or Enterprise subscription the bill is per seat: this is the API-equivalent cost.")
    header(us, 4, ["Model", "Calls", "Input tokens", "Cache write 5m tokens", "Cache write 1h tokens", "Cache read tokens",
                   "Output tokens", "Cost (USD)", f"Cost ({ccy})"])
    detail = q("Claude Usage Detail")
    r = 4
    for m in claude_models(close):
        r += 1
        us.cell(row=r, column=1, value=m)
        for c, col in zip(range(2, 10), "IJKLMNOP"):
            # "="&A forces an exact match: a bare model name such as <synthetic> would read as a comparison
            put(us, f"{chr(64 + c)}{r}", f'=SUMIFS({detail}!${col}$5:${col}${last_detail},{detail}!$E$5:$E${last_detail},"="&$A{r})',
                MONEY if c >= 8 else TOKENS)
    add_table(us, "tblClaudeUsage", "A", 4, "I", r)
    r += 1
    put(us, f"A{r}", "Total")
    for c in range(2, 10):
        col = chr(64 + c)
        put(us, f"{col}{r}", f"=SUM({col}5:{col}{r - 1})", MONEY if c >= 8 else TOKENS)
    k["total_row"](us, r, 9)
    widths(us, [28, 10, 14, 16, 16, 16, 14, 14, 14])

    # Project Timeline
    if close["timeline"]:
        ts = wb.create_sheet("Project Timeline")
        put(ts, "A1", "Project timeline", font=k["TITLE"])
        project = next((t for t in close["timeline"] if t["category"] == "project"), {})
        put(ts, "A2", f"Started {project.get('started', '')}; "
                      + (f"completed {project['completed']}" if project.get("completed") else "still running")
                      + f"; {project.get('days', '')} calendar days, {project.get('detail', '')}. "
                      "From git history of the BMad artifacts and sprint status, the agents' sanctums and the token ledger.")
        header(ts, 4, ["Category", "Name", "Detail", "Status", "Started", "Completed", "Days", "Source"])
        r = 4
        for t in close["timeline"]:
            r += 1
            values = [t["category"], t["name"], t.get("detail", ""), t.get("status", ""), iso_datetime(t.get("started", "")),
                      iso_datetime(t.get("completed", "")), int(t["days"]) if str(t.get("days", "")).isdigit() else None,
                      t.get("source", "")]
            for c, v in enumerate(values, 1):
                cell = ts.cell(row=r, column=c, value=v)
                if c in (5, 6):
                    cell.number_format = "yyyy-mm-dd"
        add_table(ts, "tblTimeline", "A", 4, "H", r)
        widths(ts, [12, 36, 30, 12, 12, 12, 8, 14])
        ts.freeze_panes = "A5"


def write_facts(model: dict, figures: dict, out: Path) -> None:
    """The cost lines as values, one row each, for dashboards that read files (Power BI, Excel Online)."""
    ccy, name = model["currency"], model["name"]
    rows = [["estimate", "currency", "category", "item", "monthly", "delivery", "run", "total"]]
    for e in figures["environments"]:
        rows.append([name, ccy, "environment", e["name"], e["monthly"], e["delivery"], e["run"], round(e["delivery"] + e["run"], 2)])
    rows.append([name, ccy, "team", "Delivery team", "", figures["team"]["cost"], 0, figures["team"]["cost"]])
    rows.append([name, ccy, "other", "Other one-off costs", "", figures["other_costs"], 0, figures["other_costs"]])
    claude = figures.get("claude")
    if claude:
        rows.append([name, ccy, "claude" if claude["in_totals"] else "claude (memo)", "Claude API usage", "",
                     claude["cost"], 0, claude["cost"]])
    rows.append([name, ccy, "contingency", "Contingency", "", "", "", figures["contingency"]])
    g = figures["grand_total"]
    rows.append([name, ccy, "grand total", "Grand total", "", g["delivery"], g["run"], g["total"]])
    with out.open("w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(rows)


# --------------------------------------------------------------------------- main


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    p.add_argument("project_root")
    p.add_argument("model")
    p.add_argument("--rate-card", type=Path)
    p.add_argument("--out", type=Path)
    p.add_argument("--check", action="store_true", help="validate only")
    p.add_argument("--drift", action="store_true", help="report cited files changed since the build")
    p.add_argument("--force", action="store_true", help="overwrite a workbook edited after the last build")
    args = p.parse_args(sys.argv[1:] if argv is None else argv)

    root = Path(args.project_root).resolve()
    model_path = Path(args.model)
    if not model_path.is_absolute():
        model_path = (root / model_path) if (root / model_path).is_file() else estimates_folder(root) / model_path
    try:
        model = json.loads(model_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e:
        print(json.dumps({"status": "error", "message": f"cannot read model {model_path}: {e}"}))
        return 2

    if args.drift:
        result = drift(root, model)
        print(json.dumps({"model": str(model_path), **result}, indent=1))
        return 1 if result["status"] == "drift" else 0

    card, card_path = load_rate_card(root, model, args.rate_card)
    errors = validate(root, model, card)
    close, close_errors = load_close_out(root, model)
    errors += close_errors
    if model.get("team") and not card:
        errors.append(f"no rate card: write {card_path} (roles and day rates) or add 'rate_card' to the model")
    if errors:
        print(json.dumps({"status": "invalid", "model": str(model_path), "errors": errors}, indent=1))
        return 1
    figures = compute(model, card, close)
    if args.check:
        print(json.dumps({"status": "valid", "model": str(model_path), "figures": figures}, indent=1))
        return 0

    out = args.out or model_path.with_name(model_path.name.removesuffix(".estimate.json").removesuffix(".json") + ".xlsx")
    if not out.resolve().is_relative_to(root):
        print(json.dumps({"status": "error", "message": f"workbook must be inside the project: {out}"}))
        return 2
    saved = model.get("stamp", {}).get("workbook_mtime")
    if out.is_file() and not args.force:
        if saved is None or abs(out.stat().st_mtime - saved) > 1:
            print(json.dumps({"status": "error", "workbook": str(out),
                              "message": "the workbook was changed after the last build (by hand?); rebuilding "
                                         "loses those changes. Carry them into the model, or pass --force"}))
            return 2
    stamp = make_stamp(root, model, card_path, close["files"] if close else [])
    try:
        sheets = write_workbook(model, card, stamp, out, close)
        facts = out.with_name(out.stem + ".facts.csv")
        write_facts(model, figures, facts)
    except PermissionError:
        print(json.dumps({"status": "error", "message": f"cannot write {out}: is it open in Excel?"}))
        return 2
    model["stamp"] = {**stamp, "workbook_mtime": out.stat().st_mtime}
    model_path.write_text(json.dumps(model, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"status": "ok", "workbook": str(out), "facts": str(facts), "model": str(model_path), "sheets": sheets,
                      "figures": figures}, indent=1))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
